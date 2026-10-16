"""Strict CSV/XLSX reader and schema mapping. No candidate names hardcoded."""
from __future__ import annotations
import csv
import io
import re
from pathlib import Path
from .common import DataError, integer, key, norm, teryt


def read_table(path, sheet=None, encoding=None, delimiter=None):
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet is None and len(book.sheetnames) != 1:
            book.close()
            raise DataError('Skoroszyt ma wiele arkuszy: podaj sheet w konfiguracji.')
        try:
            values = list(book[sheet or book.sheetnames[0]].values)
        finally:
            book.close()
    else:
        raw = path.read_bytes()
        if encoding:
            text = raw.decode(encoding)
        else:
            try:
                text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = raw.decode('cp1250')
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(text[:20000], delimiters=';,\t').delimiter
            except csv.Error as exc:
                raise DataError('Nie rozpoznano separatora CSV: podaj delimiter.') from exc
        values = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not values:
        raise DataError(f'Pusta tabela: {path}')
    headers = [str(x or '').strip().lstrip('\ufeff') for x in values[0]]
    if any(not h for h in headers) or len(set(headers)) != len(headers):
        raise DataError('Puste lub powtórzone nagłówki tabeli.')
    rows = []
    for lineno, row in enumerate(values[1:], 2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        if len(row) != len(headers):
            raise DataError(f'Wiersz {lineno}: {len(row)} wartości zamiast {len(headers)}.')
        rows.append(dict(zip(headers, row)))
    return headers, rows


def column(headers, explicit, aliases=(), prefix=None, required=True):
    if explicit:
        if explicit not in headers:
            raise DataError(f'Brak pola {explicit!r}')
        return explicit
    hits = [h for h in headers if norm(h) in aliases or (prefix and norm(h).startswith(prefix))]
    if len(hits) == 1:
        return hits[0]
    if not required and not hits:
        return None
    raise DataError(f'Nie można jednoznacznie ustalić pola {aliases or prefix}: {hits}. Podaj mapowanie.')


def prepare(path, city, mapping=None):
    mapping = mapping or {}
    headers, source = read_table(path, mapping.get('sheet'), mapping.get('encoding'), mapping.get('delimiter'))
    nf = column(headers, mapping.get('number'), ('nr komisji','nr obwodu','numer','numer obwodu','numer obwodu glosowania'))
    tf = column(headers, mapping.get('teryt'), ('teryt gminy','teryt','kod teryt','kod teryt gminy'))
    kf = column(headers, mapping.get('kind'), ('typ obwodu',))
    vf = column(headers, mapping.get('valid'), prefix='liczba glosow waznych oddanych lacznie')
    invalid = column(headers, mapping.get('invalid'), prefix='liczba glosow niewaznych', required=False)
    # PKW places candidate fields after the valid-votes total. Require the surname
    # in uppercase and at least one given name; ambiguous exports need explicit mapping.
    candidates = mapping.get('candidates')
    if candidates is None:
        candidates = {}
        for h in headers[headers.index(vf)+1:]:
            words = h.replace('_', ' ').split()
            if len(words) < 2 or not words[0].isupper():
                raise DataError(f'Nieznana kolumna po sumie głosów: {h}. Podaj candidates.')
            candidate = re.sub('[^A-Z0-9_]', '', norm(words[0]).upper().replace(' ', '_'))
            if candidate in candidates:
                candidate += '_' + norm(' '.join(words[1:])).upper().replace(' ', '_')
            if candidate in candidates:
                raise DataError('Niejednoznaczne identyfikatory kandydatów.')
            candidates[candidate] = h
    if len(candidates) < 2:
        raise DataError('Wymagane co najmniej dwie kolumny kandydatów.')
    for name, field in candidates.items():
        if not re.fullmatch('[A-Z][A-Z0-9_]{0,39}', name) or field not in headers:
            raise DataError(f'Niepoprawne mapowanie kandydata: {name}: {field}')
    if len(set(candidates.values())) != len(candidates):
        raise DataError('Ta sama kolumna wskazana dla kilku kandydatów.')
    output, seen = [], set()
    for row in source:
        if teryt(row[tf]) != teryt(city):
            continue
        identity = key(city, row[nf])
        if identity in seen:
            raise DataError(f'Powtórzony klucz {identity}')
        seen.add(identity)
        votes = {c: integer(row[f], f'{identity}/{f}') for c, f in candidates.items()}
        valid = integer(row[vf], f'{identity}/głosy ważne')
        if sum(votes.values()) != valid:
            raise DataError(f'{identity}: suma głosów kandydatów {sum(votes.values())} != {valid}')
        if not str(row[kf] or '').strip():
            raise DataError(f'{identity}: pusty typ obwodu')
        output.append(dict(ID_OBWODU=identity, NR_KOMISJI=integer(row[nf],minimum=1),
            TERYT=teryt(city), TYP_OBWODU=str(row[kf]), STALY=norm(row[kf])=='staly',
            GLOSY_WAZNE=valid, GLOSY_NIEWAZNE=integer(row[invalid]) if invalid else None,
            votes=votes))
    if not output:
        raise DataError(f'Brak protokołów dla TERYT {city}')
    return dict(rows=output, candidates=candidates, headers=headers)


def vote_fields(row, turn):
    valid, votes = row['GLOSY_WAZNE'], row['votes']
    prefix = f'T{turn}_'
    result = {prefix+'GLOSY_WAZNE': valid, prefix+'GLOSY_NIEWAZNE': row['GLOSY_NIEWAZNE']}
    for candidate, count in votes.items():
        result[prefix+'GLOSY_'+candidate] = count
        result[prefix+'PCT_'+candidate] = count*100/valid if valid else None
    maximum = max(votes.values())
    leaders = sorted(c for c in votes if votes[c] == maximum) if valid else []
    leader = leaders[0] if len(leaders)==1 else ('REMIS' if leaders else 'BRAK_GLOSOW')
    pct = maximum*100/valid if valid else None
    category = sum(pct >= b for b in (30,40,50,60)) if pct is not None else -1
    result.update({prefix+'LIDER':leader, prefix+'LIDERZY':'|'.join(leaders),
        prefix+'LIDER_PCT':pct, prefix+'KLASA':category, prefix+'STYL':f'{leader}_{category}'})
    if len(votes)==2 and valid:
        result[prefix+'PRZEWAGA_PP'] = abs(max(votes.values())-min(votes.values()))*100/valid
    return result


def join(boundaries, rounds):
    regular = [{r['ID_OBWODU']:r for r in table['rows'] if r['STALY']} for table in rounds]
    ids = {b['properties']['ID_OBWODU'] for b in boundaries}
    if len(ids) != len(boundaries):
        raise DataError('Powtórzone identyfikatory poligonów.')
    differences = [{"brak_geometrii": sorted(set(table)-ids), "brak_wynikow": sorted(ids-set(table))} for table in regular]
    if any(d['brak_geometrii'] or d['brak_wynikow'] for d in differences):
        raise DataError(f'Niezgodność granic i protokołów stałych: {differences}')
    output = []
    for feature in boundaries:
        props = dict(feature['properties'])
        for turn, table in enumerate(regular,1):
            props.update(vote_fields(table[props['ID_OBWODU']],turn))
        output.append(dict(geometry=feature['geometry'],properties=props))
    return output
